# -*- coding: utf-8 -*-
"""读取 DSVA国标对照表.xlsx 的完整内容"""
import sys
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\j1877\WorkBuddy\2026-08-26-15-05-33\DSVA国标对照表.xlsx', data_only=True)
ws = wb['DSVA国标对照表']

print('=== 完整内容 ===\n')
for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    # 过滤全空行
    if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
        continue
    print(f'【第 {row_idx} 行】')
    for col_idx, cell in enumerate(row, start=1):
        col_letter = chr(64 + col_idx) if col_idx <= 26 else 'A' + chr(64 + col_idx - 26)
        if cell is not None:
            cell_str = str(cell)
            if len(cell_str) > 100:
                cell_str = cell_str[:100] + '...'
            print(f'  {col_letter}: {cell_str}')
    print()
