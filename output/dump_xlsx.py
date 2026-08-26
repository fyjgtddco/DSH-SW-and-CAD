# -*- coding: utf-8 -*-
"""读取全部 xlsx 内容，输出到文本文件便于分析"""
import sys
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\j1877\WorkBuddy\2026-08-26-15-05-33\DSVA国标对照表.xlsx', data_only=True)
ws = wb['DSVA国标对照表']

with open(r'C:\Users\j1877\Desktop\dsh-engineering-mode\output\xlsx_dump.txt', 'w', encoding='utf-8') as out:
    out.write(f'总行数: {ws.max_row}\n')
    out.write(f'总列数: {ws.max_column}\n\n')
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
            continue
        out.write(f'=== 第 {row_idx} 行 ===\n')
        for col_idx, cell in enumerate(row, start=1):
            if cell is not None:
                cell_str = str(cell)
                out.write(f'  列{col_idx}: {cell_str}\n')
        out.write('\n')

print('已写入 xlsx_dump.txt')
print(f'行数: {ws.max_row}, 列数: {ws.max_column}')
